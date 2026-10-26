"""Evidence-linked reviewer response and manuscript correction draft, not submission."""
import argparse
import json
from pathlib import Path


def main():
    p=argparse.ArgumentParser();p.add_argument('--run',type=Path,required=True)
    p.add_argument('--grid',required=True);args=p.parse_args()
    summary=json.loads((args.run/args.grid/'summary.json').read_text())
    decision=json.loads((args.run/args.grid/'decision.json').read_text())
    status=json.loads((args.run/'status.json').read_text())
    points=[
        ('Novelty / neuro-symbolic','Explicit supervisor predicates, states and transitions specified. Contribution remains application/method evaluation; no first-of-kind claim. Broader verified literature comparison remains incomplete.'),
        ('Latency reward','New causal PPO pilots omit exogenous latency penalty and privileged counterfactual bonus. Executed tracking/heading/effort/smoothness determine reward. This corrects the causal claim; it is not a matched retraining reproduction of every original reward ablation.'),
        ('Delay implementation','One-in-flight request with original observation timestamps; plant evolves each10ms while unavailable. No pending output/forecast/service oracle enters policy. Simulated-service and actual worker tests are distinct.'),
        ('LQR delay contradiction','PPO service delay is separated from local sensor/controller latency. Long PPO stalls do not mean delayed LQR feedback is tolerated. Actual backup response timing reported separately.'),
        ('49x speedup','Original speedup withdrawn. Inference and command-processing distributions are measured;100Hz is not used as PPO inference duration. No cross-hardware speed ranking.'),
        ('Ablation timing','Each evaluated arm records actual inference/processing time. Reward/gain removal is not assigned a100Hz inference frequency. Asynchronous processing costs remain separate.'),
        ('LQR degradation','LQR-only, same-supervisor constants, causal supervised initialization, scratch and fine-tuned PPO included. Negative gains and heading/steering tradeoffs retained; no automatic promotion.'),
        ('Generalization','Development uses known sine/DLC families with exact mirrored physical/noise conditions. New sine parameters reserved only for conditional confirmation. This is not broad unseen-road generalization.'),
        ('Stability','Nominal LQR Lyapunov result limited to its linear assumptions. No zero convergence from bounded residual, nonlinear global safety or guaranteed delay robustness is claimed. Saturation/actuator constraints are experimental checks, not proof.'),
        ('Missing details','Frozen protocol includes reward coefficients, categorical5 distribution, preprocessing, seeds, budgets, gains/recovery and delay handling. Three training seeds, raw traces, source/model provenance and episode-level metrics retained.'),
        ('Energy / blind distance','Unsupported GPU energy and end-to-end freshness comparisons removed. No energy measurements or physical sensor-to-actuator latency are claimed.'),
        ('Numerics / citations','New tables generated from raw traces and independently recomputed. Original bibliography is not certified; primary metadata and complete citation verification remain pending.')]
    lines=['# Reviewer response: experimental evidence draft','',
        'This document records work performed and remaining limitations. It is not a claim that all reviewer concerns are resolved, nor a submitted response.','']
    for i,(title,text) in enumerate(points,1):lines += [f'## Comment {i}: {title}','',text,'']
    lines+=['## Current evidence','',f'Grid: {args.grid}; episodes: {decision["episodes"]}; eligibility: {decision["eligible"]}.',
            '',f'See [{args.grid} results]({args.grid}/README.md) and status.json for external/uncompleted work.','',
            '## Uncompleted external/theoretical requirements','']
    lines += [f'- **{key}:** {value}' for key,value in status['external'].items()]
    (args.run/'reviewer-response.md').write_text('\n'.join(lines)+'\n')
    paper=['# Latency-Aware Residual Vehicle Control with Predictive Runtime Supervision','',
        '**Revised research draft; not publication-ready.** Numerical results below belong only to the identified campaign.','',
        '## Abstract','',
        'We evaluate bounded PPO residual assistance around a strong LQR baseline using a causal asynchronous execution interface and an explicit predictive supervisor. The supervisor combines packet freshness/validity, nominal cost comparison and recovery persistence. Matched constants and supervised initialization distinguish learning from filtering. Simulated-service tracking and actual process timing are reported separately. Formal safety, energy advantage and real-vehicle readiness are not established.','',
        '## Method','',
        'The command is clip(delta_LQR + w * residual, -0.5, 0.5). The worker returns a categorical proposal in {-0.08,-0.04,0,0.04,0.08}rad. Policy inputs exclude pending outputs, future completion and configured service duration. The local observer/LQR/supervisor executes independently of the worker; independence is a design requiring timing validation. The nominal predictor compares baseline and candidate commands over200ms, with a50ms candidate hold followed by LQR. A50ms observation-age cutoff rejects stale proposals; candidate weights are0.1/0.25/0.5 with0.05/tick increases and3 good control ticks for admission. Exact protocol.json defines reward, actuator, costs and settings.','',
        '## Evaluation','',
        'Six25k PPO pilots: three paired seeds, scratch versus causal supervised initialization; pretraining is extra. All episodes log true lateral/heading error, actuator movement, timing and rejection reason. The simple latency-gain arm is a disclosed reconstruction, not a published-controller reproduction.','',
        '| Method | Lateral RMSE (m) | Gain vs LQR | Survival |','|---|---:|---:|---:|']
    for name,m in summary.items():paper.append(f'| {name} | {m["rmse_m"]:.6f} | {m["gain_vs_lqr_pct"]:.3f}% | {m["survived"]}/{m["episodes"]} |')
    paper += ['', '## Interpretation and limits','',
        f'Development eligibility: {decision["eligible"]}. All failed criteria are retained in decision.json. Known-path fresh noise is not unseen-road generalization. Whole training/noise pairs, rather than control ticks, are the independent comparison blocks. No arbitrary real-time, stability, GPU speedup or energy guarantee is inferred.', '',
        '## References','',
        'Primary uploaded literature includes Wabersich and Zeilinger predictive safety filtering; Tearle et al. learning-based racing safety filtering; Cano Cordoba et al. delayed-observation shielding; Kalaria et al. delay-aware robust driving/racing; and Tabrizian et al. delay-aware merging. Exact publication metadata and a broader runtime-assurance review must be verified before submission.']
    (args.run/'revised-manuscript-draft.md').write_text('\n'.join(paper)+'\n')
    try:
        from docx import Document
        doc=Document()
        for line in paper:
            if line.startswith('# '):doc.add_heading(line[2:],0)
            elif line.startswith('## '):doc.add_heading(line[3:],1)
            elif line:doc.add_paragraph(line)
        doc.save(args.run/'Revised_LANS_Evidence_Draft.docx')
    except ImportError:
        (args.run/'docx-unavailable.txt').write_text('python-docx missing; Markdown draft preserved.\n')


if __name__=='__main__':main()
